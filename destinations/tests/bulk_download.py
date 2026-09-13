import csv
import io
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from destinations.api.v1.admin.serializers import (
    AdminActivityBulkUploadSerializer,
    BULK_ACTIVITY_TEMPLATE,
    BULK_DESTINATION_TEMPLATE,
)
from destinations.models import Activity, ActivityImage, Attraction, Cuisine, Destination


class DestinationBulkDownloadAPITests(TestCase):
    url = "/api/v1/admin/destinations/bulk-download/"

    def setUp(self):
        self.user = User.objects.create_superuser(
            email="bulk-download@example.com",
            password="password",
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.destination = Destination.objects.create(
            name="Pokhara",
            country="Nepal",
            country_code="NPL",
            region="Gandaki",
            destination_type="city",
            latitude=28.2096,
            longitude=83.9856,
            tagline="Lakeside city",
            description="Gateway to the Annapurna region.",
            cover_image="https://example.com/pokhara.jpg",
            budget_tier="mid",
            difficulty_level="easy",
            local_languages=["Nepali", "English"],
            best_travel_months=[10, 11],
            currency="Nepalese Rupee",
            currency_code="NPR",
            notes=["Carry cash"],
            picking_reasons=["Mountain views"],
        )
        self.attraction = Attraction.objects.create(
            destination=self.destination,
            name="Phewa Lake",
            attraction_type="natural_site",
            description="A scenic freshwater lake.",
            best_months=[10, 11],
        )
        self.activity = Activity.objects.create(
            destination=self.destination,
            name="Paragliding",
            activity_type="adventure",
            description="Tandem paragliding over the valley.",
            budget_tier="premium",
            best_months=[9, 10, 11],
            booking_required=True,
        )
        self.cuisine = Cuisine.objects.create(
            destination=self.destination,
            name="Thakali Set",
            description="Traditional rice meal.",
            meal_type="lunch",
        )

    def test_default_xlsx_contains_exact_destination_template_sheets_and_rows(self):
        response = self.client.get(
            self.url,
            {"destination_ids": str(self.destination.id)},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(".xlsx", response["Content-Disposition"])
        workbook = load_workbook(io.BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(
            workbook.sheetnames,
            list(BULK_DESTINATION_TEMPLATE["xlsx_sheets"]),
        )
        for sheet_name, expected_headers in BULK_DESTINATION_TEMPLATE["xlsx_sheets"].items():
            headers = [cell.value for cell in next(workbook[sheet_name].iter_rows())]
            self.assertEqual(headers, expected_headers)

        destination_row = list(workbook["destinations"].iter_rows(values_only=True))[1]
        destination_values = dict(
            zip(BULK_DESTINATION_TEMPLATE["xlsx_sheets"]["destinations"], destination_row)
        )
        self.assertEqual(destination_values["destination_key"], self.destination.slug)
        self.assertEqual(destination_values["local_languages"], "Nepali;English")
        self.assertEqual(workbook["attractions"].max_row, 2)
        self.assertEqual(workbook["activities"].max_row, 2)
        self.assertEqual(workbook["cuisines"].max_row, 2)

    def test_activity_csv_can_be_selected_by_id_and_revalidated_for_upload(self):
        other_activity = Activity.objects.create(
            destination=self.destination,
            name="Boating",
            activity_type="city_tour",
            description="A calm lake ride.",
            budget_tier="mid",
        )

        response = self.client.get(
            self.url,
            {
                "type": "activities",
                "format": "csv",
                "activity_ids": str(self.activity.id),
            },
        )

        self.assertEqual(response.status_code, 200)
        text = response.content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        self.assertEqual(reader.fieldnames, BULK_ACTIVITY_TEMPLATE["csv_columns"])
        rows = list(reader)
        self.assertEqual([row["name"] for row in rows], [self.activity.name])
        self.assertNotEqual(rows[0]["name"], other_activity.name)
        self.assertEqual(rows[0]["best_months"], "9;10;11")
        self.assertEqual(rows[0]["booking_required"], "true")

        target = Destination.objects.create(
            name="Kathmandu",
            country="Nepal",
            country_code="NPL",
            destination_type="city",
            latitude=27.7172,
            longitude=85.324,
            tagline="Historic capital",
            description="A cultural destination.",
            cover_image="https://example.com/kathmandu.jpg",
            budget_tier="mid",
            currency="Nepalese Rupee",
            currency_code="NPR",
        )
        upload = SimpleUploadedFile(
            "activities.csv",
            response.content,
            content_type="text/csv",
        )
        serializer = AdminActivityBulkUploadSerializer(
            data={"file": upload},
            context={"request": type("Request", (), {"user": self.user})(), "destination": target},
        )
        self.assertTrue(serializer.is_valid(), serializer.errors)

    def test_rejects_invalid_type_format_and_ids(self):
        cases = (
            ({"type": "hotels"}, "type"),
            ({"format": "pdf"}, "format"),
            ({"activity_ids": "not-a-uuid"}, "activity_ids"),
        )
        for params, expected_field in cases:
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertEqual(response.status_code, 400)
                self.assertIn(expected_field, response.data)

    @patch("destinations.api.v1.admin.views.process_vector_operations.delay")
    def test_batch_train_uses_the_same_type_and_id_parameters(self, delay_mock):
        response = self.client.post(
            f"/api/v1/admin/destinations/batch-train/?type=activities"
            f"&activity_ids={self.activity.id}"
        )

        self.assertEqual(response.status_code, 202)
        self.assertEqual(response.data["data"]["queued_ids"], [str(self.activity.id)])
        delay_mock.assert_called_once_with(
            [
                {
                    "action": "index",
                    "source_type": "activity",
                    "source_id": str(self.activity.id),
                    "destination_id": str(self.destination.id),
                }
            ]
        )

    @patch("destinations.api.v1.admin.views.process_vector_operations.delay")
    def test_default_batch_train_queues_destination_trees(self, delay_mock):
        response = self.client.post("/api/v1/admin/destinations/batch-train/")

        self.assertEqual(response.status_code, 202)
        delay_mock.assert_called_once_with(
            [
                {
                    "action": "index_tree",
                    "source_type": "destination",
                    "source_id": str(self.destination.id),
                    "destination_id": str(self.destination.id),
                }
            ]
        )

    @patch("destinations.api.v1.admin.views.delete_image")
    def test_batch_delete_requires_ids_and_deletes_only_selected_items(self, delete_mock):
        missing_ids_response = self.client.delete(
            "/api/v1/admin/destinations/batch-delete/?type=activities"
        )
        self.assertEqual(missing_ids_response.status_code, 400)
        self.assertIn("activity_ids", missing_ids_response.data)

        other_activity = Activity.objects.create(
            destination=self.destination,
            name="City Tour",
            activity_type="city_tour",
            description="A guided city tour.",
            budget_tier="mid",
        )
        ActivityImage.objects.create(
            activity=self.activity,
            image_url="https://example.com/paragliding-gallery.jpg",
        )
        self.activity.cover_image = "https://example.com/paragliding-cover.jpg"
        self.activity.save(update_fields=["cover_image", "updated_at"])

        response = self.client.delete(
            f"/api/v1/admin/destinations/batch-delete/?type=activities"
            f"&activity_ids={self.activity.id}"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["data"]["deleted_ids"], [str(self.activity.id)])
        self.assertFalse(Activity.objects.filter(id=self.activity.id).exists())
        self.assertTrue(Activity.objects.filter(id=other_activity.id).exists())
        self.assertCountEqual(
            [call.kwargs["image_url"] for call in delete_mock.call_args_list],
            [
                "https://example.com/paragliding-cover.jpg",
                "https://example.com/paragliding-gallery.jpg",
            ],
        )
